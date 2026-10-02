# -*- coding: utf-8 -*-
"""
Persian <-> Gregorian Date Converter
 - GUI mode : double-click PersianDateConverter.exe (no arguments)
 - CLI mode : called by the Word macro with argv[1]=date, argv[2]=optional format

Author  : Saeed Majidi G.
Company : Translation Office 1324
Email   : rassamtranslation@gmail.com
Version : 2.0
"""

import os
import re
import sys
import threading
import datetime as _dt
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import jdatetime
from docx import Document


# ===========================================================================
# 1. PATHS  (works both running from source and from a PyInstaller .exe)
# ===========================================================================
if getattr(sys, 'frozen', False):
    exe_dir = os.path.dirname(sys.executable)
    # When the CLI is built with --onedir, its exe lives one folder
    # deeper than the GUI's exe. Climb one level so both share the
    # same settings.ini and log.txt (next to the GUI's exe).
    if os.path.basename(sys.executable).lower().startswith(
            'persiandateconvertercli'):
        APP_DIR = os.path.dirname(exe_dir)
    else:
        APP_DIR = exe_dir
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

SETTINGS_FILE = os.path.join(APP_DIR, 'settings.ini')
LOG_FILE      = os.path.join(APP_DIR, 'log.txt')

APP_NAME    = 'Persian Date Converter'
APP_VERSION = '2.0'
APP_AUTHOR  = 'Saeed Majidi'
APP_EMAIL   = 'rassamtranslation@gmail.com'
APP_OFFICE  = 'Translation Office 1324'


# ===========================================================================
# 2. LOGGER
# ===========================================================================
def log_event(kind, message):
    try:
        stamp = _dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        with open(LOG_FILE, 'a', encoding='utf-8') as f:
            f.write(f'[{stamp}] [{kind}] {message}\n')
    except Exception:
        pass


# ===========================================================================
# 3. SETTINGS FILE
# ===========================================================================
def load_settings():
    data = {}
    if os.path.isfile(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    if not line or line.startswith('[') or '=' not in line:
                        continue
                    k, v = line.split('=', 1)
                    data[k.strip()] = v.strip()
        except Exception:
            pass
    return data


def save_setting(key, value):
    data = load_settings()
    data[key] = str(value)
    try:
        with open(SETTINGS_FILE, 'w', encoding='utf-8') as f:
            f.write('[general]\n')
            for k, v in data.items():
                f.write(f'{k} = {v}\n')
    except Exception as e:
        log_event('ERROR', f'save_setting failed: {e}')


# ===========================================================================
# 4. MONTH NAMES (multi-language)
# ===========================================================================
MONTHS = {
    'en': {
        'full': ['January','February','March','April','May','June',
                 'July','August','September','October','November','December'],
        'abbr': ['Jan.','Feb.','Mar.','Apr.','May','Jun.',
                 'Jul.','Aug.','Sep.','Oct.','Nov.','Dec.'],
    },
    'fr': {
        'full': ['janvier','février','mars','avril','mai','juin',
                 'juillet','août','septembre','octobre','novembre','décembre'],
        'abbr': ['janv.','févr.','mars','avr.','mai','juin',
                 'juil.','août','sept.','oct.','nov.','déc.'],
    },
    'de': {
        'full': ['Januar','Februar','März','April','Mai','Juni',
                 'Juli','August','September','Oktober','November','Dezember'],
        'abbr': ['Jan.','Feb.','März','Apr.','Mai','Juni',
                 'Juli','Aug.','Sep.','Okt.','Nov.','Dez.'],
    },
    'it': {
        'full': ['gennaio','febbraio','marzo','aprile','maggio','giugno',
                 'luglio','agosto','settembre','ottobre','novembre','dicembre'],
        'abbr': ['gen.','feb.','mar.','apr.','mag.','giu.',
                 'lug.','ago.','set.','ott.','nov.','dic.'],
    },
    'es': {
        'full': ['enero','febrero','marzo','abril','mayo','junio',
                 'julio','agosto','septiembre','octubre','noviembre','diciembre'],
        'abbr': ['ene.','feb.','mar.','abr.','may.','jun.',
                 'jul.','ago.','sep.','oct.','nov.','dic.'],
    },
    'tr': {
        'full': ['Ocak','Şubat','Mart','Nisan','Mayıs','Haziran',
                 'Temmuz','Ağustos','Eylül','Ekim','Kasım','Aralık'],
        'abbr': ['Oca.','Şub.','Mar.','Nis.','May.','Haz.',
                 'Tem.','Ağu.','Eyl.','Eki.','Kas.','Ara.'],
    },
    'ar': {
        'full': ['يناير','فبراير','مارس','أبريل','مايو','يونيو',
                 'يوليو','أغسطس','سبتمبر','أكتوبر','نوفمبر','ديسمبر'],
        'abbr': ['يناير','فبراير','مارس','أبريل','مايو','يونيو',
                 'يوليو','أغسطس','سبتمبر','أكتوبر','نوفمبر','ديسمبر'],
    },

    'ru': {
        'full': ['января','февраля','марта','апреля','мая','июня',
                 'июля','августа','сентября','октября','ноября','декабря'],
        'abbr': ['янв.','фев.','мар.','апр.','мая','июн.',
                 'июл.','авг.','сен.','окт.','ноя.','дек.'],
    },
}


# ===========================================================================
# 5. FORMATS
# ===========================================================================
DATE_FORMATS = {
    '1':  ('{d} {B} {Y}',       'English — British full: 27 September 2024',     'en'),
    '2':  ('{m}/{d}/{Y}',       'English — American slash: 09/27/2024',           'en'),
    '3':  ('{d}-{m}-{Y}',       'English — European dash: 27-09-2024',            'en'),
    '4':  ('{Y}-{m}-{d}',       'English — ISO: 2024-09-27',                      'en'),
    '5':  ('{B} {j}, {Y}',      'English — American full: September 27, 2024',    'en'),
    '6':  ('{d} {b} {Y}',       'English — British abbrev.: 27 Sep. 2024',        'en'),
    '7':  ('{b} {j}, {Y}',      'English — American abbrev.: Sep. 27, 2024',      'en'),
    '8':  ('{j} {B} {Y}',       'French — 27 septembre 2024',                     'fr'),
    '9':  ('{j} {b} {Y}',       'French — 27 sept. 2024',                         'fr'),
    '10': ('{j}. {B} {Y}',      'German — 27. September 2024',                    'de'),
    '11': ('{j}. {b} {Y}',      'German — 27. Sep. 2024',                         'de'),
    '12': ('{j} {B} {Y}',       'Italian — 27 settembre 2024',                    'it'),
    '13': ('{j} {b} {Y}',       'Italian — 27 set. 2024',                         'it'),
    '14': ('{j} de {B} de {Y}', 'Spanish — 27 de septiembre de 2024',             'es'),
    '15': ('{j} {b} {Y}',       'Spanish — 27 sep. 2024',                         'es'),
    '16': ('{j} {B} {Y}',       'Turkish — 27 Eylül 2024',                        'tr'),
    '17': ('{j} {b} {Y}',       'Turkish — 27 Eyl. 2024',                         'tr'),
    '18': ('{j} {B} {Y}',       'Arabic — 27 سبتمبر 2024',                     'ar'),
    '19': ('{j}/{m}/{Y}',       'Arabic — 27/09/2024',                            'ar'),
    '20': ('{d}.{m}.{Y}',       'Dot — day.month.year: 27.09.2024',               'en'),
    '21': ('{Y}.{m}.{d}',       'Dot — year.month.day (ISO dot): 2024.09.27',     'en'),
    '22': ('{m}.{d}.{Y}',       'Dot — American dot: 09.27.2024',                 'en'),
    '23': ('{j}. {B} {Y}',      'Dot — day. Month year: 27. September 2024',      'en'),
    '24': ('{j} {B} {Y}',       'Russian — 27 сентября 2024',                     'ru'),
    '25': ('{j} {b} {Y}',       'Russian — 27 сен. 2024',                         'ru'),
}
DEFAULT_FORMAT_KEY = '5'


# ===========================================================================
# 6. CORE CONVERSION
# ===========================================================================
def persian_to_gregorian(p_year, p_month, p_day):
    jd = jdatetime.date(int(p_year), int(p_month), int(p_day))
    g = jd.togregorian()
    return _dt.date(g.year, g.month, g.day)


def gregorian_to_persian(g_date):
    jd = jdatetime.date.fromgregorian(date=g_date)
    return jd.year, jd.month, jd.day


def _fmt(template, g, lang):
    return (template
            .replace('{B}', MONTHS[lang]['full'][g.month - 1])
            .replace('{b}', MONTHS[lang]['abbr'][g.month - 1])
            .replace('{Y}', f'{g.year:04d}')
            .replace('{y}', f'{g.year % 100:02d}')
            .replace('{m}', f'{g.month:02d}')
            .replace('{d}', f'{g.day:02d}')
            .replace('{j}', str(g.day)))


def format_gregorian(g_date, format_key):
    if format_key not in DATE_FORMATS:
        format_key = DEFAULT_FORMAT_KEY
    template, _desc, lang = DATE_FORMATS[format_key]
    return _fmt(template, g_date, lang)


def convert_persian_string(persian_date, format_key):
    try:
        y, m, d = map(int, persian_date.split('/'))
        g = persian_to_gregorian(y, m, d)
        return format_gregorian(g, format_key)
    except Exception:
        return None


# ===========================================================================
# 7. VALIDATORS
# ===========================================================================
def validate_persian(y, m, d):
    if not (1 <= m <= 12):
        return False, 'Persian month must be between 1 and 12.'
    if m <= 6:
        max_day = 31
    elif m <= 11:
        max_day = 30
    else:
        try:
            max_day = 30 if jdatetime.date(y, 1, 1).isleap() else 29
        except Exception:
            max_day = 29
    if not (1 <= d <= max_day):
        return False, f'Persian day must be between 1 and {max_day}.'
    if not (1200 <= y <= 1600):
        return False, 'Persian year must be between 1200 and 1600.'
    return True, ''


def validate_gregorian(y, m, d):
    if not (1 <= m <= 12):
        return False, 'Gregorian month must be between 1 and 12.'
    if not (1 <= d <= 31):
        return False, 'Gregorian day must be between 1 and 31.'
    if not (1800 <= y <= 2200):
        return False, 'Gregorian year must be between 1800 and 2200.'
    try:
        _dt.datetime(y, m, d)
    except ValueError:
        return False, 'That day does not exist in that month.'
    return True, ''


# ===========================================================================
# 8. BATCH .docx
# ===========================================================================
PERSIAN_DATE_PATTERN = r'\b(\d{4})/(\d{1,2})/(\d{1,2})\b'


def replace_persian_dates(text, format_key):
    def _sub(m):
        c = convert_persian_string(m.group(0), format_key)
        return c if c else m.group(0)
    return re.sub(PERSIAN_DATE_PATTERN, _sub, text)


def _process_run(run, format_key):
    original = run.text
    updated = replace_persian_dates(original, format_key)
    if original != updated:
        run.text = updated
        return True
    return False


def process_docx_file(file_path, format_key, output_folder, log=print):
    try:
        doc = Document(file_path)
        changed = False
        for para in doc.paragraphs:
            for run in para.runs:
                if _process_run(run, format_key):
                    changed = True
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if _process_run(run, format_key):
                                changed = True
        for section in doc.sections:
            for part in (section.header, section.footer):
                if part:
                    for para in part.paragraphs:
                        for run in para.runs:
                            if _process_run(run, format_key):
                                changed = True
        if changed:
            os.makedirs(output_folder, exist_ok=True)
            base = os.path.basename(file_path)
            new_name = base.replace('.docx', '_updated.docx')
            doc.save(os.path.join(output_folder, new_name))
            log(f'  [OK] {base} -> {new_name}')
        else:
            log(f'  [--] {os.path.basename(file_path)}: no Persian dates found.')
    except Exception as e:
        log(f'  [!!] {os.path.basename(file_path)}: {e}')


def process_folder(folder_path, format_key, output_folder, log=print):
    log(f'Scanning: {folder_path}')
    log(f'Output  : {output_folder}')
    log(f'Format  : {format_key} -> {DATE_FORMATS[format_key][1]}\n')
    count = 0
    for filename in sorted(os.listdir(folder_path)):
        if filename.lower().endswith('.docx') and not filename.startswith('~$'):
            process_docx_file(os.path.join(folder_path, filename),
                              format_key, output_folder, log)
            count += 1
    log(f'\nDone. {count} file(s) processed.')
    return count


# ===========================================================================
# 9. VBA MACRO GENERATOR
# ===========================================================================
VBA_TEMPLATE = r''''
' -------------------------------------------------------------------
'  DEFAULT - uses whatever format is set in the GUI (settings.ini)
' -------------------------------------------------------------------
Sub ConvertPersianDate()
    ConvertDateWithFormat ""
End Sub


' -------------------------------------------------------------------
'  PRIORITY FORMATS
' -------------------------------------------------------------------

' 1) American full: September 17, 2024
Sub ConvertPersianDate_Full()
    ConvertDateWithFormat "5"
End Sub

' 2) American abbreviated: Sep. 17, 2024
Sub ConvertPersianDate_Abbrev()
    ConvertDateWithFormat "7"
End Sub

' 3) British full: 17 September 2024
Sub ConvertPersianDate_British()
    ConvertDateWithFormat "1"
End Sub

' 4) British abbreviated: 17 Sep. 2024
Sub ConvertPersianDate_BritishAbbrev()
    ConvertDateWithFormat "6"
End Sub

' 5) ISO dashes: 2024-09-17
Sub ConvertPersianDate_ISO()
    ConvertDateWithFormat "4"
End Sub

' 6) ISO dots: 2024.09.17
Sub ConvertPersianDate_DotISO()
    ConvertDateWithFormat "21"
End Sub

' 7) American slashes: 09/17/2024
Sub ConvertPersianDate_SlashUS()
    ConvertDateWithFormat "2"
End Sub


' ===================================================================
'  SHARED ENGINE
' ===================================================================
Private Sub ConvertDateWithFormat(fmtKey As String)
    Dim dateInput As String
    Dim exePath As String
    Dim vbsPath As String
    Dim tmpFile As String
    Dim Q As String
    Dim cmdLine As String
    Dim convertedDate As String
    Dim wsh As Object
    Dim fso As Object
    Dim stream As Object

    Q = Chr(34)

    exePath = "{exe_path}"
    vbsPath = "{vbs_path}"
    tmpFile = Environ("TEMP") & "\pdc_word_out.txt"

    dateInput = InputBox("Enter the Persian date (YYYY/MM/DD):", _
                         "Persian Date Conversion")
    If dateInput = "" Then Exit Sub

    Set fso = CreateObject("Scripting.FileSystemObject")
    If fso.FileExists(tmpFile) Then fso.DeleteFile tmpFile, True

    If fmtKey = "" Then
        cmdLine = "cmd.exe /c " & Q & _
                  "wscript.exe " & Q & vbsPath & Q & " " & _
                  Q & exePath & Q & " " & _
                  Q & dateInput & Q & _
                  " > " & Q & tmpFile & Q & " 2>&1" & Q
    Else
        cmdLine = "cmd.exe /c " & Q & _
                  "wscript.exe " & Q & vbsPath & Q & " " & _
                  Q & exePath & Q & " " & _
                  Q & dateInput & Q & " " & Q & fmtKey & Q & _
                  " > " & Q & tmpFile & Q & " 2>&1" & Q
    End If

    Set wsh = CreateObject("WScript.Shell")
    wsh.Run cmdLine, 0, True

    If fso.FileExists(tmpFile) Then
        Set stream = CreateObject("ADODB.Stream")
        stream.Type = 2
        stream.Charset = "utf-8"
        stream.Open
        stream.LoadFromFile tmpFile
        convertedDate = Trim(stream.ReadText(-1))
        stream.Close
        Set stream = Nothing
        fso.DeleteFile tmpFile, True
    Else
        convertedDate = ""
    End If

    convertedDate = Replace(convertedDate, vbCrLf, "")
    convertedDate = Replace(convertedDate, vbLf, "")
    convertedDate = Replace(convertedDate, vbCr, "")

    If convertedDate = "" Or InStr(convertedDate, "Error:") > 0 Then
        MsgBox "Error: Please enter a valid Persian date in the format YYYY/MM/DD."
    Else
        Selection.TypeText convertedDate
    End If
End Sub
'''


def build_macro(exe_path):
    # exe_path = ...\dist\PersianDateConverterCLI\PersianDateConverterCLI.exe
    # vbs_path = ...\dist\RunHidden.vbs   (one folder up)
    exe_dir  = os.path.dirname(exe_path)
    dist_dir = os.path.dirname(exe_dir)
    vbs_path = os.path.join(dist_dir, 'RunHidden.vbs')
    return VBA_TEMPLATE.format(exe_path=exe_path, vbs_path=vbs_path)


# ===========================================================================
# 10. GUI
# ===========================================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f'{APP_NAME} V {APP_VERSION}')
        self.geometry('840x700')
        self.minsize(780, 640)

        saved = load_settings().get('format_key', DEFAULT_FORMAT_KEY)
        if saved not in DATE_FORMATS:
            saved = DEFAULT_FORMAT_KEY
        self.format_key    = tk.StringVar(value=saved)
        self.exe_path      = tk.StringVar(
            value=os.path.join(APP_DIR, 'PersianDateConverterCLI',
                               'PersianDateConverterCLI.exe'))
        self.folder_path   = tk.StringVar(value=os.getcwd())
        self.output_folder = tk.StringVar(value=os.path.join(os.getcwd(), 'UP_DATED'))

        self.p_year = tk.StringVar(); self.p_month = tk.StringVar(); self.p_day = tk.StringVar()
        self.g_year = tk.StringVar(); self.g_month = tk.StringVar(); self.g_day = tk.StringVar()
        self.english_result = tk.StringVar()
        self._updating = False          # guard to prevent handler ping-pong

        self._build_ui()
        self._wire_live_conversion()
        log_event('INFO', f'GUI started. Settings file: {SETTINGS_FILE}')

    # ------------------------------------------------------------------ UI
    def _build_ui(self):
        nb = ttk.Notebook(self); nb.pack(fill='both', expand=True, padx=8, pady=8)
        self.tab_convert = ttk.Frame(nb)
        self.tab_batch   = ttk.Frame(nb)
        self.tab_macro   = ttk.Frame(nb)
        self.tab_log     = ttk.Frame(nb)
        self.tab_about   = ttk.Frame(nb)

        nb.add(self.tab_convert, text='Single Date')
        nb.add(self.tab_batch,   text='Word Files (Batch)')
        nb.add(self.tab_macro,   text='Word Macro')
        nb.add(self.tab_log,     text='Log')
        nb.add(self.tab_about,   text='About')

        self._build_convert_tab()
        self._build_batch_tab()
        self._build_macro_tab()
        self._build_log_tab()
        self._build_about_tab()

        self.status = tk.StringVar(value='Ready.')
        ttk.Label(self, textvariable=self.status, relief='sunken',
                  anchor='w').pack(fill='x', side='bottom')

    def _build_convert_tab(self):
        f = self.tab_convert

        top = ttk.LabelFrame(f, text='Format'); top.pack(fill='x', padx=8, pady=8)
        for i, (k, (fmt, desc, lang)) in enumerate(DATE_FORMATS.items()):
            ttk.Radiobutton(top, text=f'{k}: {desc}',
                            variable=self.format_key, value=k,
                            command=self._refresh_conversion
                            ).grid(row=i // 3, column=i % 3, sticky='w',
                                   padx=6, pady=2)

        p_frame = ttk.LabelFrame(f, text='Persian date (Shamsi)')
        p_frame.pack(fill='x', padx=8, pady=4)
        ttk.Label(p_frame, text='Year').grid(row=0, column=0, padx=4, pady=4)
        self.p_year_entry = ttk.Entry(p_frame, textvariable=self.p_year, width=8)
        self.p_year_entry.grid(row=0, column=1)
        ttk.Label(p_frame, text='Month').grid(row=0, column=2, padx=4)
        self.p_month_entry = ttk.Entry(p_frame, textvariable=self.p_month, width=5)
        self.p_month_entry.grid(row=0, column=3)
        ttk.Label(p_frame, text='Day').grid(row=0, column=4, padx=4)
        self.p_day_entry = ttk.Entry(p_frame, textvariable=self.p_day, width=5)
        self.p_day_entry.grid(row=0, column=5)
        ttk.Button(p_frame, text='Today',
                   command=self._set_today_persian).grid(row=0, column=6, padx=8)

        g_frame = ttk.LabelFrame(f, text='Gregorian date')
        g_frame.pack(fill='x', padx=8, pady=4)
        ttk.Label(g_frame, text='Year').grid(row=0, column=0, padx=4, pady=4)
        self.g_year_entry = ttk.Entry(g_frame, textvariable=self.g_year, width=8)
        self.g_year_entry.grid(row=0, column=1)
        ttk.Label(g_frame, text='Month').grid(row=0, column=2, padx=4)
        self.g_month_entry = ttk.Entry(g_frame, textvariable=self.g_month, width=5)
        self.g_month_entry.grid(row=0, column=3)
        ttk.Label(g_frame, text='Day').grid(row=0, column=4, padx=4)
        self.g_day_entry = ttk.Entry(g_frame, textvariable=self.g_day, width=5)
        self.g_day_entry.grid(row=0, column=5)
        ttk.Button(g_frame, text='Today',
                   command=self._set_today_gregorian).grid(row=0, column=6, padx=8)

        self.warn_label = ttk.Label(f, text='', foreground='#b00020',
                                    font=('Segoe UI', 10, 'bold'))
        self.warn_label.pack(fill='x', padx=10, pady=(2, 0))

        r_frame = ttk.LabelFrame(f, text='Target-language result')
        r_frame.pack(fill='x', padx=8, pady=8)
        ttk.Entry(r_frame, textvariable=self.english_result,
                  font=('Segoe UI', 14), state='readonly').pack(
            fill='x', padx=6, pady=6)
        btns = ttk.Frame(r_frame); btns.pack(fill='x', padx=6, pady=(0, 6))
        ttk.Button(btns, text='Copy to Clipboard',
                   command=self._copy_result).pack(side='left')
        ttk.Button(btns, text='Clear',
                   command=self._clear_all).pack(side='left', padx=6)

    def _build_batch_tab(self):
        f = self.tab_batch
        row = ttk.Frame(f); row.pack(fill='x', padx=8, pady=8)
        ttk.Label(row, text='Input folder:').pack(side='left')
        ttk.Entry(row, textvariable=self.folder_path).pack(side='left', fill='x',
                                                           expand=True, padx=6)
        ttk.Button(row, text='Browse...',
                   command=self._browse_input).pack(side='left')

        row2 = ttk.Frame(f); row2.pack(fill='x', padx=8, pady=4)
        ttk.Label(row2, text='Output folder:').pack(side='left')
        ttk.Entry(row2, textvariable=self.output_folder).pack(side='left', fill='x',
                                                              expand=True, padx=6)
        ttk.Button(row2, text='Browse...',
                   command=self._browse_output).pack(side='left')

        row3 = ttk.Frame(f); row3.pack(fill='x', padx=8, pady=4)
        ttk.Label(row3, text='Format:').pack(side='left')
        self.batch_format_cb = ttk.Combobox(
            row3, state='readonly', width=56,
            values=[f'{k}: {v[1]}' for k, v in DATE_FORMATS.items()])
        self.batch_format_cb.current(int(DEFAULT_FORMAT_KEY) - 1)
        self.batch_format_cb.pack(side='left', padx=6)

        ttk.Button(f, text='Convert Word Files',
                   command=self._run_batch).pack(pady=8)
        self.batch_log = tk.Text(f, height=18, wrap='word')
        self.batch_log.pack(fill='both', expand=True, padx=8, pady=(0, 8))

    def _build_macro_tab(self):
        f = self.tab_macro
        ttk.Label(f, text='1. Path to PersianDateConverterCLI.exe:'
                  ).pack(anchor='w', padx=8, pady=(8, 0))
        row = ttk.Frame(f); row.pack(fill='x', padx=8)
        ttk.Entry(row, textvariable=self.exe_path).pack(
            side='left', fill='x', expand=True)
        ttk.Button(row, text='Browse...',
                   command=self._browse_exe).pack(side='left', padx=4)

        ttk.Label(f, text='2. The Word macro uses the format currently '
                          'selected in the Single Date tab (settings.ini).'
                  ).pack(anchor='w', padx=8, pady=(8, 0))

        ttk.Button(f, text='Refresh macro',
                   command=self._refresh_macro).pack(pady=8)
        self.macro_text = tk.Text(f, height=22, wrap='none', font=('Consolas', 10))
        self.macro_text.pack(fill='both', expand=True, padx=8, pady=(0, 8))
        ttk.Button(f, text='Copy macro to clipboard',
                   command=self._copy_macro).pack(pady=(0, 8))
        self._refresh_macro()

    def _build_log_tab(self):
        f = self.tab_log
        bar = ttk.Frame(f); bar.pack(fill='x', padx=8, pady=8)
        ttk.Button(bar, text='Refresh', command=self._refresh_log).pack(side='left')
        ttk.Button(bar, text='Open log file',
                   command=self._open_log_file).pack(side='left', padx=6)
        ttk.Button(bar, text='Clear log file',
                   command=self._clear_log_file).pack(side='left', padx=6)
        self.log_text = tk.Text(f, height=24, wrap='word', font=('Consolas', 9))
        self.log_text.pack(fill='both', expand=True, padx=8, pady=(0, 8))
        self._refresh_log()

    def _build_about_tab(self):
        f = self.tab_about

        ttk.Label(f, text=APP_NAME,
                  font=('Segoe UI', 20, 'bold')).pack(pady=(28, 2))

        ttk.Label(f, text=f'Version {APP_VERSION}',
                  font=('Segoe UI', 10), foreground='#666').pack()

        ttk.Separator(f, orient='horizontal').pack(fill='x', padx=60, pady=20)

        info = ttk.Frame(f)
        info.pack(pady=8)

        ttk.Label(info, text='Company',
                  font=('Segoe UI', 11)).grid(row=0, column=0,
                                              sticky='e', padx=(0, 12), pady=5)
        ttk.Label(info, text=APP_OFFICE,
                  font=('Segoe UI', 12, 'bold')).grid(row=0, column=1,
                                                      sticky='w', pady=5)

        ttk.Label(info, text='Developed by',
                  font=('Segoe UI', 11)).grid(row=1, column=0,
                                              sticky='e', padx=(0, 12), pady=5)
        ttk.Label(info, text=APP_AUTHOR,
                  font=('Segoe UI', 12, 'bold')).grid(row=1, column=1,
                                                      sticky='w', pady=5)

        ttk.Label(info, text='Email',
                  font=('Segoe UI', 11)).grid(row=2, column=0,
                                              sticky='e', padx=(0, 12), pady=5)
        email_lbl = ttk.Label(info, text=APP_EMAIL,
                              font=('Segoe UI', 11),
                              foreground='#0645AD', cursor='hand2')
        email_lbl.grid(row=2, column=1, sticky='w', pady=5)
        email_lbl.bind('<Button-1>', lambda e: self._open_mail())

        ttk.Separator(f, orient='horizontal').pack(fill='x', padx=60, pady=20)

        desc = ('Convert Persian (Shamsi) dates to Gregorian\n'
                'in 25 formats across 8 languages.\n\n'
                'Batch convert .docx files, use the clipboard,\n'
                'or press Ctrl+Shift+D or your custom shortcut inside Word.')
        ttk.Label(f, text=desc, justify='center',
                  font=('Segoe UI', 10)).pack(pady=6)

        ttk.Separator(f, orient='horizontal').pack(fill='x', padx=60, pady=20)

        btns = ttk.Frame(f)
        btns.pack(pady=8)

        ttk.Button(btns, text='Send Email',
                   command=self._open_mail).pack(side='left', padx=4)
        ttk.Button(btns, text='Copy Email',
                   command=self._copy_email).pack(side='left', padx=4)
        ttk.Button(btns, text='Open Log Folder',
                   command=self._open_log_folder).pack(side='left', padx=4)

        ttk.Label(f, text=f'© 2026 {APP_AUTHOR} — {APP_OFFICE}',
                  font=('Segoe UI', 9),
                  foreground='#888').pack(side='bottom', pady=18)

    # ------------------------------------------------------------ helpers
    def _browse_input(self):
        d = filedialog.askdirectory(initialdir=self.folder_path.get() or os.getcwd())
        if d: self.folder_path.set(d)

    def _browse_output(self):
        d = filedialog.askdirectory(initialdir=self.output_folder.get() or os.getcwd())
        if d: self.output_folder.set(d)

    def _browse_exe(self):
        f = filedialog.askopenfilename(title='Select PersianDateConverterCLI.exe',
                                       filetypes=[('Executable', '*.exe'),
                                                  ('All files', '*.*')])
        if f:
            self.exe_path.set(f); self._refresh_macro()

    def _copy_result(self):
        v = self.english_result.get()
        if v:
            self.clipboard_clear(); self.clipboard_append(v)
            self.status.set(f'Copied: {v}')
            log_event('INFO', f'Copied: {v}')

    def _copy_macro(self):
        self.clipboard_clear()
        self.clipboard_append(self.macro_text.get('1.0', 'end-1c'))
        self.status.set('Macro copied to clipboard.')

    def _clear_all(self):
        for v in (self.p_year, self.p_month, self.p_day,
                  self.g_year, self.g_month, self.g_day):
            v.set('')
        self.english_result.set('')
        self._clear_field_error()

    def _set_today_persian(self):
        jd = jdatetime.date.today()
        self.p_year.set(str(jd.year)); self.p_month.set(str(jd.month)); self.p_day.set(str(jd.day))

    def _set_today_gregorian(self):
        gd = _dt.date.today()
        self.g_year.set(str(gd.year)); self.g_month.set(str(gd.month)); self.g_day.set(str(gd.day))

    def _refresh_log(self):
        self.log_text.delete('1.0', 'end')
        if os.path.isfile(LOG_FILE):
            with open(LOG_FILE, 'r', encoding='utf-8', errors='replace') as fh:
                self.log_text.insert('end', fh.read())
        self.log_text.see('end')

    def _open_log_file(self):
        if os.path.isfile(LOG_FILE):
            os.startfile(LOG_FILE)
        else:
            messagebox.showinfo('Log', 'No log file yet.')

    def _clear_log_file(self):
        if messagebox.askyesno('Confirm', 'Delete the current log file?'):
            try:
                if os.path.isfile(LOG_FILE):
                    os.remove(LOG_FILE)
                self._refresh_log()
                log_event('INFO', 'Log cleared by user.')
            except Exception as e:
                messagebox.showerror('Error', str(e))

    def _show_field_error(self, message):
        self.warn_label.config(text='\u26a0  ' + message)
        for w in (self.p_year_entry, self.p_month_entry, self.p_day_entry,
                  self.g_year_entry, self.g_month_entry, self.g_day_entry):
            w.configure(foreground='#b00020')
        self.english_result.set('')
        self.status.set('Invalid input: ' + message)

    def _clear_field_error(self):
        if self.warn_label.cget('text'):
            self.warn_label.config(text='')
        for w in (self.p_year_entry, self.p_month_entry, self.p_day_entry,
                  self.g_year_entry, self.g_month_entry, self.g_day_entry):
            w.configure(foreground='')

    # -------------------------------------------------- About tab helpers
    def _open_mail(self):
        import webbrowser
        webbrowser.open(f'mailto:{APP_EMAIL}')

    def _copy_email(self):
        self.clipboard_clear()
        self.clipboard_append(APP_EMAIL)
        self.status.set('Email copied to clipboard.')

    def _open_log_folder(self):
        try:
            os.startfile(APP_DIR)
        except Exception as e:
            messagebox.showerror('Error', str(e))

    # -------------------------------------------------- live conversion
    def _wire_live_conversion(self):
        for var in (self.p_year, self.p_month, self.p_day):
            var.trace_add('write', lambda *_: self._persian_changed())
        for var in (self.g_year, self.g_month, self.g_day):
            var.trace_add('write', lambda *_: self._gregorian_changed())
        self.format_key.trace_add('write', lambda *_: self._refresh_conversion())
        self.format_key.trace_add(
            'write',
            lambda *_: (save_setting('format_key', self.format_key.get()),
                        log_event('INFO',
                                  f'Format changed to {self.format_key.get()} '
                                  f'({DATE_FORMATS[self.format_key.get()][1]})'))
        )

    def _persian_changed(self):
        if self._updating:
            return
        y_s, m_s, d_s = self.p_year.get(), self.p_month.get(), self.p_day.get()
        if not (y_s and m_s and d_s):
            self._clear_field_error(); return
        try:
            y, m, d = int(y_s), int(m_s), int(d_s)
        except ValueError:
            self._show_field_error('Persian year/month/day must be numbers.'); return
        ok, err = validate_persian(y, m, d)
        if not ok:
            self._show_field_error(err); return
        self._clear_field_error()
        try:
            g = persian_to_gregorian(y, m, d)
        except Exception as e:
            self._show_field_error(f'Conversion failed: {e}'); return
        self._updating = True
        try:
            self.g_year.set(str(g.year))
            self.g_month.set(str(g.month))
            self.g_day.set(str(g.day))
        finally:
            self._updating = False
        result = format_gregorian(g, self.format_key.get())
        self.english_result.set(result)
        self.status.set('Converted from Persian.')
        log_event('PERSIAN->GREG',
                  f'{y}/{m}/{d} -> {g.isoformat()} '
                  f'(format {self.format_key.get()} -> {result})')

    def _gregorian_changed(self):
        if self._updating:
            return
        y_s, m_s, d_s = self.g_year.get(), self.g_month.get(), self.g_day.get()
        if not (y_s and m_s and d_s):
            self._clear_field_error(); return
        try:
            y, m, d = int(y_s), int(m_s), int(d_s)
        except ValueError:
            self._show_field_error('Gregorian year/month/day must be numbers.'); return
        ok, err = validate_gregorian(y, m, d)
        if not ok:
            self._show_field_error(err); return
        self._clear_field_error()
        try:
            gd = _dt.date(y, m, d)
            jy, jm, jd = gregorian_to_persian(gd)
        except Exception as e:
            self._show_field_error(f'Conversion failed: {e}'); return
        self._updating = True
        try:
            self.p_year.set(str(jy))
            self.p_month.set(str(jm))
            self.p_day.set(str(jd))
        finally:
            self._updating = False
        result = format_gregorian(gd, self.format_key.get())
        self.english_result.set(result)
        self.status.set('Converted from Gregorian.')
        log_event('GREG->PERSIAN',
                  f'{gd.isoformat()} -> {jy}/{jm}/{jd} '
                  f'(format {self.format_key.get()} -> {result})')


    def _refresh_conversion(self):
        if self.p_year.get() and self.p_month.get() and self.p_day.get():
            self._persian_changed()
        elif self.g_year.get() and self.g_month.get() and self.g_day.get():
            self._gregorian_changed()

    def _run_batch(self):
        fmt_key = self.batch_format_cb.get().split(':', 1)[0].strip()
        folder  = self.folder_path.get()
        out     = self.output_folder.get()
        if not os.path.isdir(folder):
            messagebox.showerror('Error', f'Input folder does not exist:\n{folder}')
            log_event('ERROR', f'Batch aborted: bad folder {folder}')
            return
        self.batch_log.delete('1.0', 'end')
        self.status.set('Running batch conversion...')
        log_event('BATCH', f'Started folder={folder}, out={out}, format={fmt_key}')

        def log(msg):
            self.batch_log.insert('end', msg + '\n')
            self.batch_log.see('end')
            self.update_idletasks()
            log_event('BATCH', msg.strip())

        def worker():
            try:
                n = process_folder(folder, fmt_key, out, log)
                self.status.set(f'Batch complete. {n} file(s).')
                log_event('BATCH', f'Done. {n} file(s).')
            except Exception as e:
                log(f'FATAL: {e}')
                log_event('ERROR', f'Batch fatal: {e}')
                self.status.set('Batch failed.')

        threading.Thread(target=worker, daemon=True).start()

    def _refresh_macro(self):
        self.macro_text.delete('1.0', 'end')
        self.macro_text.insert('1.0', build_macro(self.exe_path.get()))


# ===========================================================================
# 11. CLI MODE
# ===========================================================================
def cli_mode(argv):
    # Force UTF-8 on stdout so non-ASCII output (Arabic, French accents,
    # Turkish characters, etc.) doesn't crash on Windows' default cp1252.
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

    if len(argv) < 2:
        return 1
    persian_date = argv[1]
    if len(argv) > 2:
        fmt_key = argv[2]
        source  = 'arg'
    else:
        fmt_key = load_settings().get('format_key', DEFAULT_FORMAT_KEY)
        source  = 'settings.ini'

    result = convert_persian_string(persian_date, fmt_key)
    if result:
        log_event('CLI', f'{persian_date} format={fmt_key} ({source}) -> {result}')
        try:
            print(result)
        except UnicodeEncodeError:
            sys.stdout.buffer.write(result.encode('utf-8'))
            sys.stdout.buffer.write(b'\n')
        return 0

    log_event('CLI-WARN', f'Bad date from Word: {persian_date!r} format={fmt_key}')
    try:
        print("Error: Please enter a valid Persian date in the format YYYY/MM/DD.")
    except UnicodeEncodeError:
        sys.stdout.buffer.write(
            b"Error: Please enter a valid Persian date in the format YYYY/MM/DD.\n")
    return 1


# ===========================================================================
# 12. ENTRY POINT
# ===========================================================================
def main():
    if len(sys.argv) > 1:
        sys.exit(cli_mode(sys.argv))
    App().mainloop()


if __name__ == '__main__':
    main()